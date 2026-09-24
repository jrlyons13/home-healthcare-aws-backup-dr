"""Shared PowerPoint theme: dark navy, teal/orange cards, Segoe UI."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

BG = RGBColor(0x0D, 0x16, 0x27)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEAL = RGBColor(0x48, 0xC0, 0xD8)
ORANGE = RGBColor(0xF9, 0xA8, 0x25)
MUTED = RGBColor(0x8F, 0xA3, 0xB8)

FONT = "Segoe UI"
DEFAULT_FOOTER_LEFT = "HOME HEALTHCARE AWS BACKUP/DR • TEAM REVIEW"
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _accent_color(name: str) -> RGBColor:
    return ORANGE if name == "orange" else TEAL


def _set_slide_background(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def _add_textbox(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height)


def _style_run(run, *, size: int, color: RGBColor, bold: bool = False, caps: bool = False) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    if caps:
        run.font.all_caps = True


def _add_header(slide, title: str, subtitle: str | None = None) -> None:
    title_box = _add_textbox(slide, Inches(0.65), Inches(0.45), Inches(12), Inches(0.9))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = title
    _style_run(run, size=32, color=WHITE, bold=True)

    if subtitle:
        sub_box = _add_textbox(slide, Inches(0.65), Inches(1.15), Inches(12), Inches(0.55))
        stf = sub_box.text_frame
        stf.clear()
        sp = stf.paragraphs[0]
        sp.alignment = PP_ALIGN.LEFT
        srun = sp.add_run()
        srun.text = subtitle
        _style_run(srun, size=16, color=TEAL, bold=False)


def _add_footer(slide, page: int, footer_left: str) -> None:
    left_box = _add_textbox(slide, Inches(0.65), Inches(7.05), Inches(10), Inches(0.35))
    ltf = left_box.text_frame
    ltf.clear()
    lp = ltf.paragraphs[0]
    lr = lp.add_run()
    lr.text = footer_left
    _style_run(lr, size=9, color=MUTED, caps=True)

    right_box = _add_textbox(slide, Inches(12.35), Inches(7.05), Inches(0.6), Inches(0.35))
    rtf = right_box.text_frame
    rtf.clear()
    rp = rtf.paragraphs[0]
    rp.alignment = PP_ALIGN.RIGHT
    rr = rp.add_run()
    rr.text = f"{page:02d}"
    _style_run(rr, size=11, color=MUTED, bold=True)


def _add_card(slide, left, top, width, height, label: str, body: str, accent: str) -> None:
    color = _accent_color(accent)
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.background()
    shape.line.color.rgb = color
    shape.line.width = Pt(1.25)

    label_box = _add_textbox(slide, left + Inches(0.2), top + Inches(0.15), width - Inches(0.4), Inches(0.35))
    ltf = label_box.text_frame
    ltf.clear()
    lp = ltf.paragraphs[0]
    lr = lp.add_run()
    lr.text = label
    _style_run(lr, size=10, color=color, bold=True, caps=True)

    body_box = _add_textbox(slide, left + Inches(0.2), top + Inches(0.55), width - Inches(0.4), height - Inches(0.7))
    btf = body_box.text_frame
    btf.clear()
    btf.word_wrap = True
    bp = btf.paragraphs[0]
    br = bp.add_run()
    br.text = body
    _style_run(br, size=13, color=WHITE)


def _add_cards_slide(slide, data: dict) -> None:
    _add_header(slide, data["title"], data.get("subtitle"))
    cards = data["cards"]
    count = len(cards)
    gap = Inches(0.35)
    margin = Inches(0.65)
    total_gap = gap * (count - 1)
    card_w = (SLIDE_W - (margin * 2) - total_gap) / count
    card_h = Inches(3.35)
    top = Inches(2.05)
    for i, card in enumerate(cards):
        left = margin + (card_w + gap) * i
        _add_card(slide, left, top, card_w, card_h, card["label"], card["body"], card.get("accent", "teal"))


def _add_bullets_slide(slide, data: dict) -> None:
    _add_header(slide, data["title"], data.get("subtitle"))
    box = _add_textbox(slide, Inches(0.85), Inches(2.0), Inches(11.8), Inches(4.5))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, bullet in enumerate(data["bullets"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(10)
        run = p.add_run()
        run.text = f"• {bullet}"
        _style_run(run, size=17, color=WHITE)


def _add_hero_slide(slide, data: dict) -> None:
    title_box = _add_textbox(slide, Inches(0.65), Inches(2.4), Inches(12), Inches(1.2))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = data["title"]
    _style_run(run, size=40, color=WHITE, bold=True)

    if data.get("subtitle"):
        sub_box = _add_textbox(slide, Inches(0.65), Inches(3.55), Inches(12), Inches(0.7))
        stf = sub_box.text_frame
        stf.clear()
        sp = stf.paragraphs[0]
        srun = sp.add_run()
        srun.text = data["subtitle"]
        _style_run(srun, size=18, color=TEAL)


def add_slide(prs: Presentation, data: dict, page: int, footer_left: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_background(slide)

    layout = data.get("layout", "bullets")
    if layout == "hero":
        _add_hero_slide(slide, data)
    elif layout == "cards":
        _add_cards_slide(slide, data)
    else:
        _add_bullets_slide(slide, data)

    _add_footer(slide, page, footer_left)
    slide.notes_slide.notes_text_frame.text = data.get("notes", "")


def build_presentation(
    slides: list[dict],
    output: Path,
    *,
    footer_left: str = DEFAULT_FOOTER_LEFT,
) -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    for idx, slide_data in enumerate(slides, start=1):
        add_slide(prs, slide_data, idx, footer_left)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
